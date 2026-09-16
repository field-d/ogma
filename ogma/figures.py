from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

class FigureCounter:
    def __init__(self):
        self.count = 0

    def next(self):
        self.count += 1
        return self.count

def add_figure(doc, fig_counter, image_path, title, note=None, width=Inches(6)):


    figure_number = fig_counter.next()

    # figure number
    number_para = doc.add_paragraph()
    number_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    number_run = number_para.add_run(f"Figure {figure_number}")
    number_run.bold = True

    # title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_para.add_run(title)
    title_run.italic = True

    doc.add_picture(image_path, width=width)

    last_para = doc.paragraphs[-1]
    last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if note:
        note_para = doc.add_paragraph()
        note_run = note_para.add_run(f"Note. {note}")
        note_run.italic = True
        note_run.font.size = doc.styles["Normal"].font.size

    return figure_number