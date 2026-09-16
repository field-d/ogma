from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH

def create_document():
    doc = Document()

    # margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # default font
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # paragraph formatting
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)

    return doc


def add_title_page(doc, title, name, student_num, course, university, module_code, lecturer, date):

    # pushes the title down vertically
    for _ in range(3):
        doc.add_paragraph()

    # title of the paper
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.bold = True

    # blank line
    doc.add_paragraph()

    # name
    name_para = doc.add_paragraph(name)
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # student number
    student_num_para = doc.add_paragraph(student_num)
    student_num_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # course
    course_para = doc.add_paragraph(course)
    course_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # university
    uni_para = doc.add_paragraph(university)
    uni_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # module
    module_para = doc.add_paragraph(module_code)
    module_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # lecturer
    lecturer_para = doc.add_paragraph(lecturer)
    lecturer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # date
    date_para = doc.add_paragraph(date)
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

def add_heading(doc, text, level=1):
    heading_para = doc.add_paragraph()

    if level == 1:
        heading_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading_para.add_run(text)
        run.bold = True
    elif level == 2:
        heading_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = heading_para.add_run(text)
    elif level == 3:
        heading_para_alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = heading_para.add_run(text)
        run.bold = True
        run.italic= True

    return heading_para

def add_standard_sections(doc, sections=None, page_break_between=True):
    if sections is None:
        sections = ["Abstract", "Introduction", "Method", "Results", "Discussion", "References","Appendix"]

    for i, section in enumerate(sections):
        if page_break_between and i > 0:
            doc.add_page_break()

        add_heading(doc, section, level=1)
        doc.add_paragraph()